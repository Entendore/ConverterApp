import os
import threading
import zipfile
import tarfile
import shutil
import csv
import json
import xml.etree.ElementTree as ET
import re
import configparser
import sqlite3
import base64
import pickle
import plistlib
from kivy.app import App
from kivy.uix.screenmanager import ScreenManager, Screen
from kivy.uix.popup import Popup
from kivy.uix.filechooser import FileChooserListView
from kivy.core.window import Window
from kivy.lang import Builder
from kivy.clock import Clock

# Set default window size for desktop testing
Window.size = (500, 700)

# Load KV language string
Builder.load_string('''
<MainMenuScreen>:
    BoxLayout:
        orientation: 'vertical'
        padding: 30
        spacing: 10
        
        Label:
            text: 'Universal File Converter'
            font_size: 32
            size_hint_y: 0.12
            color: 0.2, 0.6, 0.8, 1
            
        ScrollView:
            size_hint_y: 0.88
            do_scroll_x: False
            
            GridLayout:
                cols: 1
                spacing: 8
                size_hint_y: None
                height: self.minimum_height
                
                Button:
                    text: '🖼️ Image & Medical Conversion'
                    font_size: 16
                    size_hint_y: None
                    height: '48dp'
                    on_release: app.root.current = 'image'
                    
                Button:
                    text: '🎵 Audio Conversion'
                    font_size: 16
                    size_hint_y: None
                    height: '48dp'
                    on_release: app.root.current = 'audio'
                    
                Button:
                    text: '🎬 Video Conversion'
                    font_size: 16
                    size_hint_y: None
                    height: '48dp'
                    on_release: app.root.current = 'video'
                    
                Button:
                    text: '📄 Document & eBook Conversion (Inc. Apple Pages)'
                    font_size: 16
                    size_hint_y: None
                    height: '48dp'
                    on_release: app.root.current = 'text'
                    
                Button:
                    text: '🗄️ Archive Conversion'
                    font_size: 16
                    size_hint_y: None
                    height: '48dp'
                    on_release: app.root.current = 'archive'
                    
                Button:
                    text: '📊 Data, Geo & Localization (GPX, PO, ENV)'
                    font_size: 16
                    size_hint_y: None
                    height: '48dp'
                    on_release: app.root.current = 'data'
                    
                Button:
                    text: '🔧 Advanced Formats (DB, Fonts, PLIST)'
                    font_size: 16
                    size_hint_y: None
                    height: '48dp'
                    on_release: app.root.current = 'advanced'
                    
                Button:
                    text: '🛠️ Utilities & Dev Tools (QR, Base64, Logs)'
                    font_size: 16
                    size_hint_y: None
                    height: '48dp'
                    on_release: app.root.current = 'utility'
                    
                Button:
                    text: 'ℹ️ About'
                    font_size: 16
                    size_hint_y: None
                    height: '48dp'
                    on_release: app.root.current = 'about'

<ConversionScreen>:
    BoxLayout:
        orientation: 'vertical'
        padding: 20
        spacing: 15
        
        BoxLayout:
            orientation: 'horizontal'
            size_hint_y: 0.1
            
            Button:
                text: '← Back'
                size_hint_x: 0.3
                on_release: app.root.current = 'main'
                
            Label:
                text: root.title
                font_size: 18
                color: 0.2, 0.6, 0.8, 1
                
        BoxLayout:
            orientation: 'vertical'
            spacing: 15
            size_hint_y: 0.8
            
            Button:
                id: btn_select
                text: '📁 Select File'
                size_hint_y: 0.15
                on_release: app.open_file_chooser(root.file_type)
                
            Label:
                id: file_label
                text: 'No file selected'
                size_hint_y: 0.1
                color: 0.5, 0.5, 0.5, 1
                
            Spinner:
                id: format_spinner
                text: 'Select Output Format'
                values: root.formats
                size_hint_y: 0.15
                
            Button:
                id: btn_convert
                text: '🔄 Convert'
                size_hint_y: 0.15
                disabled: True
                on_release: app.start_conversion(root.file_type)
                
            ProgressBar:
                id: progress
                size_hint_y: 0.1
                disabled: True
                
            Label:
                id: status_label
                text: 'Status: Ready'
                size_hint_y: 0.15
                halign: 'center'
                valign: 'middle'
                text_size: self.size

<AboutScreen>:
    BoxLayout:
        orientation: 'vertical'
        padding: 30
        spacing: 20
        
        Label:
            text: 'Universal File Converter\\n\\nVersion 9.0\\n\\nThe ultimate cross-platform file manipulation tool!'
            font_size: 18
            halign: 'center'
            valign: 'middle'
            text_size: self.size
            
        Button:
            text: '← Back to Main Menu'
            on_release: app.root.current = 'main'

<FileChooserPopup>:
    BoxLayout:
        orientation: 'vertical'
        
        FileChooserListView:
            id: filechooser
            multiselect: False
            dirselect: False
            filters: root.filters
            
        BoxLayout:
            size_hint_y: 0.2
            spacing: 10
            
            Button:
                text: 'Cancel'
                on_release: root.dismiss()
                
            Button:
                text: 'Select'
                on_release: app.select_file(filechooser.selection, root.file_type)
''')

class MainMenuScreen(Screen):
    pass

class ConversionScreen(Screen):
    def __init__(self, **kwargs):
        super(ConversionScreen, self).__init__(**kwargs)
        self.title = ''
        self.file_type = ''
        self.formats = []

class AboutScreen(Screen):
    pass

class FileChooserPopup(Popup):
    def __init__(self, file_type, **kwargs):
        super(FileChooserPopup, self).__init__(**kwargs)
        self.file_type = file_type
        self.title = f"Select {file_type.title()} File"
        self.size_hint = (0.9, 0.9)
        
        if file_type == 'image':
            self.ids.filechooser.filters = ['*.png', '*.jpg', '*.jpeg', '*.bmp', '*.gif', '*.tiff', '*.webp', '*.ico', '*.heic', '*.heif', '*.svg', '*.psd', '*.dcm']
        elif file_type == 'audio':
            self.ids.filechooser.filters = ['*.mp3', '*.wav', '*.ogg', '*.flac', '*.aac', '*.m4a', '*.wma', '*.aiff', '*.opus', '*.amr', '*.au']
        elif file_type == 'video':
            self.ids.filechooser.filters = ['*.mp4', '*.avi', '*.mkv', '*.mov', '*.wmv', '*.flv', '*.webm', '*.3gp', '*.mpeg', '*.mpg', '*.ogv', '*.ts', '*.m2ts']
        elif file_type == 'text':
            # Added .pages
            self.ids.filechooser.filters = ['*.pdf', '*.docx', '*.txt', '*.rtf', '*.odt', '*.epub', '*.pptx', '*.pages']
        elif file_type == 'archive':
            self.ids.filechooser.filters = ['*.zip', '*.tar', '*.gz', '*.tar.gz']
        elif file_type == 'data':
            # Added .gpx, .po, .env
            self.ids.filechooser.filters = ['*.csv', '*.json', '*.xml', '*.srt', '*.vtt', '*.xlsx', '*.md', '*.markdown', '*.html', '*.htm', '*.yaml', '*.yml', '*.toml', '*.ini', '*.rst', '*.ass', '*.ssa', '*.ipynb', '*.properties', '*.kml', '*.pkl', '*.log', '*.gpx', '*.po', '*.pot', '*.env']
        elif file_type == 'advanced':
            # Added .plist
            self.ids.filechooser.filters = ['*.ttf', '*.otf', '*.ics', '*.vcf', '*.db', '*.sqlite', '*.sqlite3', '*.plist']
        elif file_type == 'utility':
            self.ids.filechooser.filters = ['*.txt', '*.b64', '*.log', '*']

class FileConverterApp(App):
    def build(self):
        self.selected_file_path = None
        self.output_dir = self.user_data_dir
        self.conversion_thread = None
        
        sm = ScreenManager()
        
        sm.add_widget(MainMenuScreen(name='main'))
        
        image_screen = ConversionScreen(name='image')
        image_screen.title = 'Image & Medical Conversion'
        image_screen.file_type = 'image'
        image_screen.formats = ['PNG', 'JPEG', 'BMP', 'GIF', 'TIFF', 'WebP', 'ICO']
        sm.add_widget(image_screen)
        
        audio_screen = ConversionScreen(name='audio')
        audio_screen.title = 'Audio Conversion'
        audio_screen.file_type = 'audio'
        audio_screen.formats = ['MP3', 'WAV', 'OGG', 'FLAC', 'AAC', 'M4A', 'AIFF', 'OPUS', 'WMA', 'AMR', 'AU']
        sm.add_widget(audio_screen)
        
        video_screen = ConversionScreen(name='video')
        video_screen.title = 'Video Conversion'
        video_screen.file_type = 'video'
        video_screen.formats = ['MP4', 'AVI', 'MKV', 'MOV', 'WMV', 'WebM', '3GP', 'FLV', 'MPEG', 'OGV', 'TS']
        sm.add_widget(video_screen)
        
        text_screen = ConversionScreen(name='text')
        text_screen.title = 'Document Conversion'
        text_screen.file_type = 'text'
        text_screen.formats = ['PDF', 'DOCX', 'TXT', 'EPUB']
        sm.add_widget(text_screen)
        
        archive_screen = ConversionScreen(name='archive')
        archive_screen.title = 'Archive Conversion'
        archive_screen.file_type = 'archive'
        archive_screen.formats = ['ZIP', 'TAR']
        sm.add_widget(archive_screen)
        
        # Data Screen Formats (Added ENV, PO)
        data_screen = ConversionScreen(name='data')
        data_screen.title = 'Data, Geo & Localization Conversion'
        data_screen.file_type = 'data'
        data_screen.formats = ['CSV', 'JSON', 'XML', 'YAML', 'XLSX', 'MD', 'HTML', 'SRT', 'VTT', 'TOML', 'INI', 'RST', 'ASS', 'PY', 'GeoJSON', 'ENV', 'PO']
        sm.add_widget(data_screen)
        
        # Advanced Screen Formats (Added JSON for PLIST)
        advanced_screen = ConversionScreen(name='advanced')
        advanced_screen.title = 'Advanced Format Conversion'
        advanced_screen.file_type = 'advanced'
        advanced_screen.formats = ['OTF', 'TTF', 'TXT', 'JSON', 'CSV']
        sm.add_widget(advanced_screen)
        
        utility_screen = ConversionScreen(name='utility')
        utility_screen.title = 'Utilities & Dev Tools'
        utility_screen.file_type = 'utility'
        utility_screen.formats = ['QR_PNG', 'B64', 'BIN', 'JSON', 'CSV']
        sm.add_widget(utility_screen)
        
        sm.add_widget(AboutScreen(name='about'))
        
        return sm
    
    def open_file_chooser(self, file_type):
        content = FileChooserPopup(file_type)
        content.open()
    
    def select_file(self, selection, file_type):
        if selection:
            self.selected_file_path = selection[0]
            filename = os.path.basename(self.selected_file_path)
            screen = self.root.get_screen(file_type)
            screen.ids.file_label.text = f'Selected: {filename}'
            screen.ids.btn_convert.disabled = False
            screen.ids.status_label.text = 'Status: File selected. Ready to convert.'
            
        for widget in App.get_running_app().root_window.children:
            if isinstance(widget, Popup):
                widget.dismiss()
    
    def start_conversion(self, file_type):
        screen = self.root.get_screen(file_type)
        target_format = screen.ids.format_spinner.text
        
        if target_format == 'Select Output Format':
            screen.ids.status_label.text = 'Status: Please select an output format.'
            return
        
        screen.ids.btn_convert.disabled = True
        screen.ids.progress.disabled = False
        screen.ids.progress.value = 0
        screen.ids.status_label.text = 'Status: Converting...'
        
        self.conversion_thread = threading.Thread(
            target=self.convert_file,
            args=(file_type, target_format)
        )
        self.conversion_thread.daemon = True
        self.conversion_thread.start()
        
        Clock.schedule_interval(self.update_progress, 0.1)
    
    def update_progress(self, dt):
        screen = self.root.current_screen
        if screen and hasattr(screen, 'ids') and 'progress' in screen.ids:
            screen.ids.progress.value += 2
            if screen.ids.progress.value >= 100:
                screen.ids.progress.value = 100
                Clock.unschedule(self.update_progress)
    
    def convert_file(self, file_type, target_format):
        try:
            if file_type == 'image':
                self.convert_image(target_format)
            elif file_type == 'audio':
                self.convert_audio(target_format)
            elif file_type == 'video':
                self.convert_video(target_format)
            elif file_type == 'text':
                self.convert_text(target_format)
            elif file_type == 'archive':
                self.convert_archive(target_format)
            elif file_type == 'data':
                self.convert_data_code(target_format)
            elif file_type == 'advanced':
                self.convert_advanced(target_format)
            elif file_type == 'utility':
                self.convert_utility(target_format)
        except Exception as e:
            Clock.schedule_once(lambda dt: self.show_error(str(e)))
        finally:
            Clock.schedule_once(lambda dt: self.conversion_complete())
    
    def convert_image(self, target_format):
        from PIL import Image
        try:
            from pillow_heif import register_heif_opener
            register_heif_opener()
        except ImportError:
            pass

        screen = self.root.get_screen('image')
        screen.ids.status_label.text = 'Status: Processing image...'
        input_ext = os.path.splitext(self.selected_file_path)[1].lower()
        
        if input_ext == '.dcm':
            import pydicom
            from pydicom.pixel_data_handlers.util import apply_voi_lut
            ds = pydicom.dcmread(self.selected_file_path)
            pixel_array = apply_voi_lut(ds.pixel_array, ds)
            pixel_array = (pixel_array - pixel_array.min()) / (pixel_array.max() - pixel_array.min()) * 255
            img = Image.fromarray(pixel_array.astype('uint8')).convert('L')
            base_name = os.path.splitext(os.path.basename(self.selected_file_path))[0]
            ext = target_format.lower().replace('jpeg', 'jpg')
            output_path = os.path.join(self.output_dir, f"{base_name}_converted.{ext}")
            img.save(output_path, format=target_format)
            Clock.schedule_once(lambda dt: self.show_success(output_path))
            return
        
        if input_ext == '.svg':
            from svglib.svglib import svg2rlg
            from reportlab.graphics import renderPM
            drawing = svg2rlg(self.selected_file_path)
            base_name = os.path.splitext(os.path.basename(self.selected_file_path))[0]
            ext = target_format.lower().replace('jpeg', 'jpg')
            output_path = os.path.join(self.output_dir, f"{base_name}_converted.{ext}")
            renderPM.drawToFile(drawing, output_path, fmt=target_format.lower().replace('jpeg', 'jpg'))
            Clock.schedule_once(lambda dt: self.show_success(output_path))
            return
        
        if input_ext == '.psd':
            from psd_tools import PSDImage
            psd = PSDImage.open(self.selected_file_path)
            img = psd.composite()
            if target_format in ['JPEG', 'BMP'] and img.mode in ('RGBA', 'P'):
                img = img.convert('RGB')
            base_name = os.path.splitext(os.path.basename(self.selected_file_path))[0]
            ext = target_format.lower().replace('jpeg', 'jpg')
            output_path = os.path.join(self.output_dir, f"{base_name}_converted.{ext}")
            img.save(output_path, format=target_format)
            Clock.schedule_once(lambda dt: self.show_success(output_path))
            return
        
        img = Image.open(self.selected_file_path)
        if target_format in ['JPEG', 'BMP'] and img.mode in ('RGBA', 'P'):
            img = img.convert('RGB')
        base_name = os.path.splitext(os.path.basename(self.selected_file_path))[0]
        ext = target_format.lower().replace('jpeg', 'jpg')
        output_path = os.path.join(self.output_dir, f"{base_name}_converted.{ext}")
        img.save(output_path, format=target_format)
        Clock.schedule_once(lambda dt: self.show_success(output_path))
    
    def convert_audio(self, target_format):
        import ffmpeg
        screen = self.root.get_screen('audio')
        screen.ids.status_label.text = 'Status: Processing audio...'
        base_name = os.path.splitext(os.path.basename(self.selected_file_path))[0]
        ext = target_format.lower()
        output_path = os.path.join(self.output_dir, f"{base_name}_converted.{ext}")
        ffmpeg.input(self.selected_file_path).output(output_path, format=ext, acodec='libmp3lame' if target_format == 'MP3' else None).run()
        Clock.schedule_once(lambda dt: self.show_success(output_path))
    
    def convert_video(self, target_format):
        import ffmpeg
        screen = self.root.get_screen('video')
        screen.ids.status_label.text = 'Status: Processing video...'
        base_name = os.path.splitext(os.path.basename(self.selected_file_path))[0]
        ext = target_format.lower()
        output_path = os.path.join(self.output_dir, f"{base_name}_converted.{ext}")
        ffmpeg.input(self.selected_file_path).output(output_path, format=ext, vcodec='libx264', acodec='aac').run()
        Clock.schedule_once(lambda dt: self.show_success(output_path))
    
    def convert_text(self, target_format):
        from pypdf import PdfReader
        from pdf2docx import Converter
        from docx import Document
        from fpdf import FPDF
        import striprtf
        import ebooklib
        from ebooklib import epub
        from bs4 import BeautifulSoup
        
        screen = self.root.get_screen('text')
        screen.ids.status_label.text = 'Status: Processing document...'
        base_name = os.path.splitext(os.path.basename(self.selected_file_path))[0]
        ext = target_format.lower()
        output_path = os.path.join(self.output_dir, f"{base_name}_converted.{ext}")
        input_ext = os.path.splitext(self.selected_file_path)[1].lower()
        
        # NEW: Apple Pages to PDF
        if input_ext == '.pages' and target_format == 'PDF':
            if not zipfile.is_zipfile(self.selected_file_path):
                raise Exception("Invalid PAGES file.")
            with zipfile.ZipFile(self.selected_file_path, 'r') as zf:
                pdf_name = None
                for name in zf.namelist():
                    if name.endswith('preview.pdf'):
                        pdf_name = name
                        break
                if not pdf_name:
                    raise Exception("No PDF preview found in PAGES file.")
                with zf.open(pdf_name) as src, open(output_path, 'wb') as dst:
                    shutil.copyfileobj(src, dst)
            Clock.schedule_once(lambda dt: self.show_success(output_path))
            return
        
        if target_format == 'TXT':
            if input_ext == '.pdf':
                reader = PdfReader(self.selected_file_path)
                text = "\n".join([page.extract_text() for page in reader.pages])
            elif input_ext == '.docx':
                doc = Document(self.selected_file_path)
                text = "\n".join([para.text for para in doc.paragraphs])
            elif input_ext == '.rtf':
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: text = striprtf.rtf_to_text(f.read())
            elif input_ext == '.epub':
                book = epub.read_epub(self.selected_file_path)
                text = ""
                for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                    soup = BeautifulSoup(item.get_content(), 'html.parser')
                    text += soup.get_text() + "\n"
            elif input_ext == '.odt':
                from odf.opendocument import load
                from odf import text as odf_text
                doc = load(self.selected_file_path)
                text = ""
                for para in doc.getElementsByType(odf_text.P): text += str(para) + "\n"
            elif input_ext == '.pptx':
                from pptx import Presentation
                prs = Presentation(self.selected_file_path)
                text = ""
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"): text += shape.text + "\n"
                    text += "\n"
            else:
                shutil.copy(self.selected_file_path, output_path)
            with open(output_path, 'w', encoding='utf-8') as f: f.write(text)
        elif target_format == 'DOCX':
            if input_ext == '.pdf':
                cv = Converter(self.selected_file_path)
                cv.convert(output_path, start=0, end=None)
                cv.close()
            else: raise Exception("Only PDF to DOCX is supported.")
        elif target_format == 'PDF':
            if input_ext == '.txt':
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: text = f.read()
            elif input_ext == '.docx':
                doc = Document(self.selected_file_path)
                text = "\n".join([para.text for para in doc.paragraphs])
            elif input_ext == '.rtf':
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: text = striprtf.rtf_to_text(f.read())
            else: raise Exception("Unsupported format for PDF conversion.")
            pdf = FPDF()
            pdf.add_page()
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.set_font("Arial", size=12)
            safe_text = text.encode('latin-1', 'replace').decode('latin-1')
            for line in safe_text.split('\n'): pdf.multi_cell(0, 10, txt=line)
            pdf.output(output_path)
        Clock.schedule_once(lambda dt: self.show_success(output_path))
    
    def convert_archive(self, target_format):
        screen = self.root.get_screen('archive')
        screen.ids.status_label.text = 'Status: Extracting and repacking...'
        base_name = os.path.splitext(os.path.basename(self.selected_file_path))[0]
        if base_name.endswith('.tar'): base_name = os.path.splitext(base_name)[0]
        ext = 'zip' if target_format == 'ZIP' else 'tar'
        output_path = os.path.join(self.output_dir, f"{base_name}_converted.{ext}")
        temp_dir = os.path.join(self.output_dir, "temp_extract")
        if os.path.exists(temp_dir): shutil.rmtree(temp_dir)
        os.makedirs(temp_dir)
        input_ext = os.path.splitext(self.selected_file_path)[1].lower()
        if input_ext == '.zip':
            with zipfile.ZipFile(self.selected_file_path, 'r') as zip_ref: zip_ref.extractall(temp_dir)
        elif input_ext == '.tar' or self.selected_file_path.endswith('.tar.gz'):
            with tarfile.open(self.selected_file_path, 'r:*') as tar_ref: tar_ref.extractall(temp_dir)
        if target_format == 'ZIP':
            with zipfile.ZipFile(output_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
                for root, dirs, files in os.walk(temp_dir):
                    for file in files:
                        fp = os.path.join(root, file)
                        zipf.write(fp, os.path.relpath(fp, temp_dir))
        elif target_format == 'TAR':
            with tarfile.open(output_path, 'w') as tarf:
                for root, dirs, files in os.walk(temp_dir):
                    for file in files:
                        fp = os.path.join(root, file)
                        tarf.add(fp, os.path.relpath(fp, temp_dir))
        shutil.rmtree(temp_dir)
        Clock.schedule_once(lambda dt: self.show_success(output_path))
    
    def convert_data_code(self, target_format):
        import yaml
        import markdown
        import html2text
        from openpyxl import Workbook, load_workbook
        
        screen = self.root.get_screen('data')
        screen.ids.status_label.text = 'Status: Processing data/code file...'
        base_name = os.path.splitext(os.path.basename(self.selected_file_path))[0]
        ext = target_format.lower()
        output_path = os.path.join(self.output_dir, f"{base_name}_converted.{ext}")
        input_ext = os.path.splitext(self.selected_file_path)[1].lower()
        
        if target_format == 'PY':
            if input_ext == '.ipynb':
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: nb = json.load(f)
                py_code = []
                for cell in nb.get('cells', []):
                    if cell.get('cell_type') == 'code': py_code.append("# %%\n" + "".join(cell.get('source', [])))
                    elif cell.get('cell_type') == 'markdown': py_code.append('# %% [markdown]\n' + "".join([f"# {l}" for l in cell.get('source', [])]))
                with open(output_path, 'w', encoding='utf-8') as f: f.write("\n\n".join(py_code))
            else: raise Exception("Only IPYNB to PY is supported.")
            
        elif target_format == 'GeoJSON':
            if input_ext == '.kml':
                tree = ET.parse(self.selected_file_path); root = tree.getroot()
                features = []
                for pm in root.findall('.//{*}Placemark'):
                    name_elem = pm.find('{*}name'); name = name_elem.text if name_elem is not None else "Unknown"
                    point_elem = pm.find('.//{*}Point')
                    if point_elem is not None:
                        coords_elem = point_elem.find('{*}coordinates')
                        if coords_elem is not None:
                            lon, lat, *_ = map(float, coords_elem.text.strip().split(','))
                            features.append({"type": "Feature", "geometry": {"type": "Point", "coordinates": [lon, lat]}, "properties": {"name": name}})
                with open(output_path, 'w', encoding='utf-8') as f: json.dump({"type": "FeatureCollection", "features": features}, f, indent=4)
            else: raise Exception("Only KML to GeoJSON is supported.")
            
        elif target_format == 'JSON' and input_ext == '.properties':
            config = {}
            with open(self.selected_file_path, 'r', encoding='utf-8') as f:
                for line in f:
                    line = line.strip()
                    if line and not line.startswith('#') and '=' in line:
                        k, v = line.split('=', 1); config[k.strip()] = v.strip()
            with open(output_path, 'w', encoding='utf-8') as f: json.dump(config, f, indent=4)
            
        elif target_format == 'CSV':
            if input_ext == '.xlsx':
                wb = load_workbook(self.selected_file_path); ws = wb.active
                with open(output_path, 'w', encoding='utf-8', newline='') as f:
                    writer = csv.writer(f)
                    for row in ws.iter_rows(): writer.writerow([cell.value for cell in row])
            elif input_ext == '.json':
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: data = json.load(f)
                if not isinstance(data, list): raise Exception("JSON must be a list of objects")
                keys = data[0].keys()
                with open(output_path, 'w', encoding='utf-8', newline='') as f:
                    writer = csv.DictWriter(f, fieldnames=keys); writer.writeheader(); writer.writerows(data)
            # NEW: GPX to CSV
            elif input_ext == '.gpx':
                tree = ET.parse(self.selected_file_path); root = tree.getroot()
                rows = []
                for trkpt in root.findall('.//{*}trkpt'):
                    lat = trkpt.get('lat')
                    lon = trkpt.get('lon')
                    ele = trkpt.find('{*}ele')
                    time = trkpt.find('{*}time')
                    rows.append({
                        'latitude': lat or '',
                        'longitude': lon or '',
                        'elevation': ele.text if ele is not None else '',
                        'time': time.text if time is not None else ''
                    })
                with open(output_path, 'w', encoding='utf-8', newline='') as f:
                    writer = csv.DictWriter(f, fieldnames=['latitude', 'longitude', 'elevation', 'time'])
                    writer.writeheader()
                    writer.writerows(rows)
            else: raise Exception("Unsupported input for CSV.")
            
        elif target_format == 'XLSX':
            if input_ext == '.csv':
                wb = Workbook(); ws = wb.active
                with open(self.selected_file_path, 'r', encoding='utf-8') as f:
                    for row in csv.reader(f): ws.append(row)
                wb.save(output_path)
            elif input_ext == '.json':
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: data = json.load(f)
                wb = Workbook(); ws = wb.active; keys = list(data[0].keys()); ws.append(keys)
                for item in data: ws.append([item[k] for k in keys])
                wb.save(output_path)
            else: raise Exception("Only CSV/JSON to XLSX is supported.")
            
        elif target_format == 'HTML':
            if input_ext in ['.md', '.markdown']:
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: md_text = f.read()
                with open(output_path, 'w', encoding='utf-8') as f: f.write(f"<html><body>{markdown.markdown(md_text)}</body></html>")
            elif input_ext == '.rst':
                from docutils import core
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: rst_text = f.read()
                with open(output_path, 'w', encoding='utf-8') as f: f.write(core.publish_string(rst_text, writer_name='html').decode('utf-8'))
            else: raise Exception("Only MD/RST to HTML is supported.")
            
        elif target_format == 'MD':
            if input_ext in ['.html', '.htm']:
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: html_text = f.read()
                with open(output_path, 'w', encoding='utf-8') as f: f.write(html2text.html2text(html_text))
            else: raise Exception("Only HTML to MD is supported.")
            
        elif target_format == 'JSON':
            if input_ext == '.csv':
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: data = list(csv.DictReader(f))
            elif input_ext == '.xml':
                tree = ET.parse(self.selected_file_path); root = tree.getroot()
                data = [{elem.tag: elem.text for elem in child} for child in root]
            elif input_ext in ['.yaml', '.yml']:
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: data = yaml.safe_load(f)
            elif input_ext in ['.toml']:
                try: import tomllib
                except ImportError: import tomli as tomllib
                with open(self.selected_file_path, 'rb') as f: data = tomllib.load(f)
            elif input_ext == '.ini':
                config = configparser.ConfigParser(); config.read(self.selected_file_path)
                data = {s: dict(config.items(s)) for s in config.sections()}
            elif input_ext == '.pkl':
                with open(self.selected_file_path, 'rb') as f: data = pickle.load(f)
            # NEW: PO to JSON
            elif input_ext in ['.po', '.pot']:
                data = {}
                current_msgid = []
                current_msgstr = []
                state = None
                with open(self.selected_file_path, 'r', encoding='utf-8') as f:
                    for line in f:
                        line = line.strip()
                        if line.startswith('msgid '):
                            if current_msgid and current_msgstr:
                                data["".join(current_msgid)] = "".join(current_msgstr)
                            current_msgid = [line[7:-1]]
                            current_msgstr = []
                            state = 'msgid'
                        elif line.startswith('msgstr '):
                            current_msgstr.append(line[8:-1])
                            state = 'msgstr'
                        elif line.startswith('"'):
                            if state == 'msgid': current_msgid.append(line[1:-1])
                            elif state == 'msgstr': current_msgstr.append(line[1:-1])
                if current_msgid and current_msgstr:
                    data["".join(current_msgid)] = "".join(current_msgstr)
            # NEW: ENV to JSON
            elif input_ext == '.env':
                data = {}
                with open(self.selected_file_path, 'r', encoding='utf-8') as f:
                    for line in f:
                        line = line.strip()
                        if line and not line.startswith('#') and '=' in line:
                            k, v = line.split('=', 1)
                            data[k.strip()] = v.strip().strip('"\'')
            else: raise Exception("Unsupported input for JSON.")
            with open(output_path, 'w', encoding='utf-8') as f: json.dump(data, f, indent=4, default=str, ensure_ascii=False)
            
        elif target_format == 'YAML':
            if input_ext == '.json':
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: data = json.load(f)
                with open(output_path, 'w', encoding='utf-8') as f: yaml.safe_dump(data, f, default_flow_style=False)
            else: raise Exception("Only JSON to YAML is supported.")
            
        elif target_format == 'TOML':
            if input_ext == '.json':
                import tomli_w
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: data = json.load(f)
                with open(output_path, 'wb') as f: tomli_w.dump(data, f)
            else: raise Exception("Only JSON to TOML is supported.")
            
        elif target_format == 'INI':
            if input_ext == '.json':
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: data = json.load(f)
                config = configparser.ConfigParser()
                for section, values in data.items(): config[section] = values
                with open(output_path, 'w', encoding='utf-8') as f: config.write(f)
            else: raise Exception("Only JSON to INI is supported.")
            
        elif target_format == 'XML':
            if input_ext == '.csv':
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: data = list(csv.DictReader(f))
            elif input_ext == '.json':
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: data = json.load(f)
            else: raise Exception("Unsupported input for XML.")
            root = ET.Element("root")
            for item in data:
                elem = ET.SubElement(root, "item")
                for k, v in item.items():
                    child = ET.SubElement(elem, k); child.text = str(v)
            tree = ET.ElementTree(root); ET.indent(tree, space="  ", level=0)
            tree.write(output_path, encoding='utf-8', xml_declaration=True)
            
        # NEW: JSON to PO
        elif target_format == 'PO':
            if input_ext == '.json':
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: data = json.load(f)
                with open(output_path, 'w', encoding='utf-8') as f:
                    f.write('msgid ""\n')
                    f.write('msgstr ""\n')
                    f.write('"Content-Type: text/plain; charset=UTF-8\\n"\n\n')
                    for k, v in data.items():
                        f.write(f'msgid "{k}"\n')
                        f.write(f'msgstr "{v}"\n\n')
            else: raise Exception("Only JSON to PO is supported.")
            
        # NEW: JSON to ENV
        elif target_format == 'ENV':
            if input_ext == '.json':
                with open(self.selected_file_path, 'r', encoding='utf-8') as f: data = json.load(f)
                with open(output_path, 'w', encoding='utf-8') as f:
                    for k, v in data.items():
                        v_str = str(v)
                        # Quote values if they contain spaces or # signs
                        if ' ' in v_str or '#' in v_str:
                            v_str = f'"{v_str}"'
                        f.write(f"{k}={v_str}\n")
            else: raise Exception("Only JSON to ENV is supported.")
            
        Clock.schedule_once(lambda dt: self.show_success(output_path))
    
    def convert_advanced(self, target_format):
        screen = self.root.get_screen('advanced')
        screen.ids.status_label.text = 'Status: Processing advanced format...'
        base_name = os.path.splitext(os.path.basename(self.selected_file_path))[0]
        ext = target_format.lower()
        output_path = os.path.join(self.output_dir, f"{base_name}_converted.{ext}")
        input_ext = os.path.splitext(self.selected_file_path)[1].lower()
        
        if input_ext in ['.db', '.sqlite', '.sqlite3'] and target_format in ['JSON', 'CSV']:
            conn = sqlite3.connect(self.selected_file_path); cursor = conn.cursor()
            cursor.execute("SELECT name FROM sqlite_master WHERE type='table';"); tables = cursor.fetchall()
            if not tables: raise Exception("No tables found in database.")
            if target_format == 'JSON':
                db_data = {}
                for table_name in tables:
                    table_name = table_name[0]
                    cursor.execute(f"SELECT * FROM {table_name}"); rows = cursor.fetchall()
                    columns = [d[0] for d in cursor.description]
                    db_data[table_name] = [dict(zip(columns, row)) for row in rows]
                with open(output_path, 'w', encoding='utf-8') as f: json.dump(db_data, f, indent=4)
            elif target_format == 'CSV':
                table_name = tables[0][0]
                cursor.execute(f"SELECT * FROM {table_name}"); rows = cursor.fetchall()
                columns = [d[0] for d in cursor.description]
                with open(output_path, 'w', encoding='utf-8', newline='') as f:
                    writer = csv.writer(f); writer.writerow(columns); writer.writerows(rows)
            conn.close()
            
        elif input_ext in ['.ttf', '.otf'] and target_format in ['TTF', 'OTF']:
            from fontTools.ttLib import TTFont
            font = TTFont(self.selected_file_path)
            if input_ext == '.ttf' and target_format == 'OTF': font.flavor = 'otf'; font.save(output_path)
            elif input_ext == '.otf' and target_format == 'TTF': font.flavor = None; font.save(output_path)
            else: shutil.copy(self.selected_file_path, output_path)
            
        elif input_ext == '.ics' and target_format == 'TXT':
            from icalendar import Calendar
            with open(self.selected_file_path, 'rb') as f: cal = Calendar.from_ical(f.read())
            text = "Calendar Events:\n" + "="*50 + "\n\n"
            for comp in cal.walk('VEVENT'):
                text += f"Event: {comp.get('summary', 'N/A')}\nStart: {comp.get('dtstart').dt if comp.get('dtstart') else 'N/A'}\nEnd: {comp.get('dtend').dt if comp.get('dtend') else 'N/A'}\n\n"
            with open(output_path, 'w', encoding='utf-8') as f: f.write(text)
            
        elif input_ext == '.vcf' and target_format == 'TXT':
            with open(self.selected_file_path, 'r', encoding='utf-8') as f: vcf = f.read()
            text = "Contacts:\n" + "="*50 + "\n\n"
            for line in vcf.split('\n'):
                if line.startswith('FN:'): text += f"Name: {line[3:]}\n"
                elif line.startswith('TEL:'): text += f"Phone: {line[4:]}\n"
                elif line.startswith('EMAIL:'): text += f"Email: {line[6:]}\n"
                elif line.startswith('END:VCARD'): text += "-"*50 + "\n\n"
            with open(output_path, 'w', encoding='utf-8') as f: f.write(text)
            
        # NEW: PLIST to JSON
        elif input_ext == '.plist' and target_format == 'JSON':
            with open(self.selected_file_path, 'rb') as f:
                data = plistlib.load(f)
            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(data, f, indent=4, default=str)
                
        else:
            raise Exception("Conversion not supported.")
        Clock.schedule_once(lambda dt: self.show_success(output_path))

    def convert_utility(self, target_format):
        screen = self.root.get_screen('utility')
        screen.ids.status_label.text = 'Status: Processing utility...'
        base_name = os.path.splitext(os.path.basename(self.selected_file_path))[0]
        ext = target_format.lower().replace('qr_png', 'png')
        output_path = os.path.join(self.output_dir, f"{base_name}_converted.{ext}")
        input_ext = os.path.splitext(self.selected_file_path)[1].lower()
        
        if target_format == 'QR_PNG':
            import qrcode
            with open(self.selected_file_path, 'r', encoding='utf-8') as f:
                text = f.read()
            qr = qrcode.QRCode(version=1, error_correction=qrcode.constants.ERROR_CORRECT_L, box_size=10, border=4)
            qr.add_data(text)
            qr.make(fit=True)
            img = qr.make_image(fill_color="black", back_color="white")
            img.save(output_path)
            
        elif target_format == 'B64':
            with open(self.selected_file_path, 'rb') as f:
                binary_data = f.read()
            b64_data = base64.b64encode(binary_data).decode('utf-8')
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(b64_data)
                
        elif target_format == 'BIN':
            with open(self.selected_file_path, 'r', encoding='utf-8') as f:
                b64_data = f.read()
            binary_data = base64.b64decode(b64_data)
            with open(output_path, 'wb') as f:
                f.write(binary_data)
                
        elif target_format in ['CSV', 'JSON'] and input_ext == '.log':
            log_pattern = r'(\S+) (\S+) (\S+) \[([\w:/]+\s[+\-]\d{4})\] "(\S+) (\S+)\s*(\S*)" (\d{3}) (\S+)'
            parsed_data = []
            with open(self.selected_file_path, 'r', encoding='utf-8', errors='ignore') as f:
                for line in f:
                    match = re.match(log_pattern, line)
                    if match:
                        parsed_data.append({
                            "ip": match.group(1),
                            "date": match.group(4),
                            "method": match.group(5),
                            "url": match.group(6),
                            "status": match.group(8)
                        })
            if target_format == 'CSV':
                with open(output_path, 'w', encoding='utf-8', newline='') as f:
                    writer = csv.DictWriter(f, fieldnames=["ip", "date", "method", "url", "status"])
                    writer.writeheader()
                    writer.writerows(parsed_data)
            elif target_format == 'JSON':
                with open(output_path, 'w', encoding='utf-8') as f:
                    json.dump(parsed_data, f, indent=4)
                    
        else:
            raise Exception("Utility conversion not supported for this file type.")
            
        Clock.schedule_once(lambda dt: self.show_success(output_path))
    
    def show_success(self, output_path):
        screen = self.root.current_screen
        screen.ids.status_label.text = f'Success! Saved to:\n{output_path}'
        screen.ids.btn_convert.disabled = False
        screen.ids.progress.disabled = True
        screen.ids.progress.value = 0
    
    def show_error(self, error_message):
        screen = self.root.current_screen
        screen.ids.status_label.text = f'Error: {error_message}'
        screen.ids.btn_convert.disabled = False
        screen.ids.progress.disabled = True
        screen.ids.progress.value = 0
    
    def conversion_complete(self):
        if self.conversion_thread and self.conversion_thread.is_alive():
            self.conversion_thread.join()

if __name__ == '__main__':
    FileConverterApp().run()